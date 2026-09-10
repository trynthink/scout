"""Propose BTB-derived cost/performance/lifetime values for meas_in.

Reads bss_meas_v2.xlsx (sheet "meas_in"), tech_crosswalk.csv (see
build_tech_crosswalk.py), and the raw BTB CSVs, and for every (measure row,
technology) pair covered by a non-`needs_review` crosswalk entry, computes
what the Energy Performance / Installed Cost / Lifetime cells for that
technology *would* be if sourced from BTB.

This never modifies bss_meas_v2.xlsx. It writes bss_meas_v2_btb_proposed.xlsx
alongside it: a copy of meas_in with proposed values substituted in (only
for the specific technology substrings that were resolved -- everything
else in a shared cell, e.g. other technologies not covered by the
crosswalk, is left byte-for-byte as-is), plus a "BTB Diff" sheet listing
every change with its old/new value and matched BTB source, for manual
review before merging anything back into the curated original.

Usage (from this directory):
    python populate_meas_in_from_btb.py
"""

import re
from pathlib import Path

import openpyxl
import pandas as pd

from unit_conversions import convert

BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "raw"
XLSX_PATH = BASE_DIR / "bss_meas_v2.xlsx"
CROSSWALK_PATH = BASE_DIR / "tech_crosswalk.csv"
OUT_PATH = BASE_DIR / "bss_meas_v2_btb_proposed.xlsx"

SHEET_NAME = "meas_in"

TIER_KEYWORDS = [
    # Longer/more specific strings first so e.g. "ESTAR" doesn't also match
    # a "Min. Efficiency ESTAR ..." style name meant for a different tier.
    ("Min. Efficiency", "Min. Efficiency"),
    ("ESTAR", "ESTAR"),
    ("Best", "Best"),
    ("Ref. Case", "Ref. Case"),
]

# canonical bound -> BTB regression-metric bound column suffix
REGRESSION_BOUND_COL = {
    "Low": "Lower Bound", "Typical": "Typical", "High": "Upper Bound"}
# canonical bound -> BTB installed-cost column suffix
COST_BOUND_COL = {"Low": "Low", "Typical": "Mid", "High": "High"}

# Technologies covered by build_tech_crosswalk.py that participate in a
# shared "heating"/"cooling"/"ventilation" Performance Units key rather
# than their own tech-name key (seen in combined-package meas_in rows,
# e.g. "(C) Ref. Case NG Boiler & Chiller"). Used only as a fallback when a
# row's Performance Units cell has no key matching the technology name
# directly.
END_USE_FALLBACK = {
    "heating": [
        "gas_boiler", "elec_boiler", "oil_boiler", "gas_furnace",
        "oil_furnace", "elec_res-heater", "rooftop_ASHP-heat",
        "comm_GSHP-heat", "pkg_terminal_HP-heat", "resistance heat",
        "furnace (NG)", "furnace (distillate)", "boiler (distillate)",
        "ASHP", "GSHP", "HPWH"],
    "cooling": [
        "gas_chiller", "centrifugal_chiller", "scroll_chiller",
        "screw_chiller", "reciprocating_chiller", "rooftop_AC",
        "rooftop_ASHP-cool", "comm_GSHP-cool", "pkg_terminal_AC-cool",
        "pkg_terminal_HP-cool", "gas_eng-driven_RTAC", "central AC",
        "room AC", "ASHP", "GSHP"],
    "ventilation": ["VAV_Vent", "CAV_Vent"],
}


def normalize_btb(df, sector):
    """Rename sector-specific columns to a common schema."""

    id_col = "Technology ID" if "Technology ID" in df.columns \
        else "Technology/Measure ID"
    rename = {id_col: "tech_id"}
    if sector == "commercial":
        rename.update({
            "Year": "Projection Year", "Scenario": "Projection Scenario"})
    return df.rename(columns=rename)


def load_btb_data():
    """Load and normalize both BTB CSVs, keyed by sector name."""

    out = {}
    for sector, fname in [
            ("residential", "btb_residential.csv"),
            ("commercial", "btb_commercial.csv")]:
        path = RAW_DIR / fname
        if not path.exists():
            raise FileNotFoundError(
                f"{path} not found -- run download_btb_data.py first.")
        out[sector] = normalize_btb(
            pd.read_csv(path, low_memory=False), sector)
    return out


def tier_for_name(name):
    """Infer a meas_in efficiency tier from a measure's Name string."""

    if not isinstance(name, str):
        return None
    for keyword, tier in TIER_KEYWORDS:
        if keyword in name:
            return tier
    return None


def techs_for_row(row):
    """Ordered, de-duplicated technology tokens referenced by a meas_in row."""

    seen = []
    for col in ["Baseline Technology", "Switched to Technology"]:
        val = row.get(col)
        if isinstance(val, str):
            for part in val.split(";"):
                part = part.strip()
                if part and part not in seen:
                    seen.append(part)
    return seen


def find_value_for_key(cell, key):
    """Return the raw (unparsed) value substring for `key` in a nested
    "key: value; key2: value2" cell, or the whole cell if it has no nested
    keys at all, or None if `key` is not present in a nested cell."""

    if not isinstance(cell, str):
        return None
    if ":" not in cell:
        return cell.strip()
    pattern = re.compile(
        re.escape(key) + r"\s*:\s*([^;]+)", re.IGNORECASE)
    match = pattern.search(cell)
    return match.group(1).strip() if match else None


def unit_for_tech(perf_units_cell, tech):
    """Resolve the meas_in performance unit that applies to `tech` in a
    given row, trying a direct tech-name key first and falling back to the
    row's heating/cooling/ventilation end-use key (see END_USE_FALLBACK)."""

    direct = find_value_for_key(perf_units_cell, tech)
    if direct:
        return direct
    if isinstance(perf_units_cell, str):
        for end_use, techs in END_USE_FALLBACK.items():
            if tech in techs:
                via_end_use = find_value_for_key(perf_units_cell, end_use)
                if via_end_use:
                    return via_end_use
    return None


def replace_value_for_key(cell, key, new_value):
    """Substitute the value for `key` in a nested cell (or replace the
    whole cell if it has no nested keys), leaving all other keys/formatting
    untouched. Returns the original cell unchanged if `key` is not found in
    a nested cell (caller should not have called this in that case)."""

    if not isinstance(cell, str):
        return cell
    if ":" not in cell:
        return new_value
    pattern = re.compile(
        r"(" + re.escape(key) + r"\s*:\s*)([^;]+)", re.IGNORECASE)
    return pattern.sub(lambda m: m.group(1) + new_value, cell, count=1)


def replace_cost_for_key(cell, key, new_cost):
    """Same as replace_value_for_key, but for the "key: new: X; key:
    existing: Y" nesting used by the Installed Cost column."""

    if not isinstance(cell, str):
        return cell
    if ":" not in cell:
        # No existing nested cost structure to preserve -- write a fresh
        # "new: X; existing: Y" pair for this (single) technology.
        return f"new: {new_cost['new']}; existing: {new_cost['existing']}"
    result = cell
    any_found = False
    for sub_key, val in new_cost.items():
        pattern = re.compile(
            r"(" + re.escape(key) + r"\s*:\s*" + re.escape(sub_key) +
            r"\s*:\s*)([^;]+)", re.IGNORECASE)
        if pattern.search(result):
            any_found = True
            result = pattern.sub(
                lambda m: m.group(1) + str(val), result, count=1)
    return result if any_found else cell


def get_btb_row(btb_data, entry):
    """Look up the single BTB row matching a crosswalk entry's technology
    ID, projection year, and projection scenario."""

    df = btb_data[entry["sector"]]
    match = df[
        (df["tech_id"] == entry["btb_technology_id"])
        & (df["Projection Year"] == entry["projection_year"])
        & (df["Projection Scenario"] == entry["projection_scenario"])]
    if len(match) != 1:
        return None
    return match.iloc[0]


def parse_currency(val):
    """Parse a BTB cost value (e.g. "$5,158" or 5158.0) into a float."""

    if pd.isna(val):
        return None
    if isinstance(val, str):
        val = val.replace("$", "").replace(",", "").strip()
        if not val:
            return None
    try:
        return float(val)
    except ValueError:
        return None


def is_zero_cost_placeholder(old_cost):
    """True if a technology's existing cost has a "new" or "existing"
    sub-value of exactly 0. A real installed cost is never $0, so this
    normally means the cost was deliberately zeroed out because the
    technology shares physical equipment with a paired heating/cooling
    technology elsewhere in the same row (e.g. a heat pump's cost entered
    once under the "heat" entry and zeroed under "cool" to avoid double-
    counting total installed cost) -- such rows are left alone rather than
    overwritten with a full BTB cost that would reintroduce that
    double-count.
    """

    if not isinstance(old_cost, str):
        return False
    for part in old_cost.split(";"):
        if ":" in part:
            val = parse_currency(part.split(":", 1)[1])
            if val == 0:
                return True
    return False


def compute_cost(btb_row, bound):
    """Return {"new": ..., "existing": ...} installed cost for a bound."""

    suffix = COST_BOUND_COL[bound]
    new_col = f"Typical New Construction Installed Cost ($2023) - {suffix}"
    existing_col = f"Typical Retrofit Installed Cost ($2023) - {suffix}"
    new_val = parse_currency(btb_row.get(new_col))
    existing_val = parse_currency(btb_row.get(existing_col))
    if new_val is None or existing_val is None:
        return None
    return {"new": round(new_val, 2), "existing": round(existing_val, 2)}


def main():
    """Build the BTB-proposed workbook and diff sheet."""

    btb_data = load_btb_data()
    crosswalk = pd.read_csv(CROSSWALK_PATH)
    crosswalk = crosswalk[~crosswalk["needs_review"].astype(bool)]
    crosswalk_by_tech_tier = {
        (row["scout_technology"], row["tier"]): row
        for _, row in crosswalk.iterrows()}

    meas_in_df = pd.read_excel(XLSX_PATH, sheet_name=SHEET_NAME)
    col_index = {name: i + 1 for i, name in enumerate(meas_in_df.columns)}

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb[SHEET_NAME]

    diff_rows = []
    for excel_row, (_, row) in enumerate(meas_in_df.iterrows(), start=2):
        tier = tier_for_name(row.get("Name"))
        if tier is None:
            continue
        # Multiple technologies in this row can share the same Energy
        # Performance/Installed Cost/Lifetime cell (e.g. a combined
        # heating+cooling package). Track each cell's contents here and
        # chain edits through it -- writing straight from the original
        # pandas row on every technology would make each technology's
        # edit clobber the previous one's, since they'd all overwrite the
        # same cell starting from its pre-edit text.
        cell_state = {
            "Energy Performance": row.get("Energy Performance"),
            "Installed Cost": row.get("Installed Cost"),
            "Lifetime": row.get("Lifetime"),
        }
        for tech in techs_for_row(row):
            entry = crosswalk_by_tech_tier.get((tech, tier))
            if entry is None:
                continue
            btb_row = get_btb_row(btb_data, entry)
            if btb_row is None:
                continue

            # -- Performance --
            perf_cell = cell_state["Energy Performance"]
            units_cell = row.get("Performance Units")
            target_unit = unit_for_tech(units_cell, tech)
            metric_idx = None
            for i in (1, 2):
                if btb_row.get(f"Regression metric {i} - Metric") == \
                        entry["btb_metric_name"]:
                    metric_idx = i
                    break
            if target_unit and metric_idx:
                raw_val = btb_row.get(
                    f"Regression metric {metric_idx} - "
                    f"{REGRESSION_BOUND_COL[entry['bound']]}")
                converted = convert(
                    entry["btb_metric_name"], target_unit, raw_val) \
                    if pd.notna(raw_val) else None
                if converted is not None:
                    old_val = find_value_for_key(perf_cell, tech)
                    new_val_str = f"{converted:.3g}"
                    if old_val != new_val_str:
                        cell_state["Energy Performance"] = \
                            replace_value_for_key(
                                perf_cell, tech, new_val_str)
                        ws.cell(
                            row=excel_row,
                            column=col_index["Energy Performance"]).value \
                            = cell_state["Energy Performance"]
                        diff_rows.append({
                            "Name": row.get("Name"), "technology": tech,
                            "column": "Energy Performance",
                            "old_value": old_val, "new_value": new_val_str,
                            "btb_technology_id": entry["btb_technology_id"],
                            "btb_display_name": entry["btb_display_name"],
                            "projection_scenario":
                                entry["projection_scenario"],
                            "projection_year": entry["projection_year"],
                        })

            # -- Installed cost --
            cost = compute_cost(btb_row, entry["bound"])
            if cost is not None:
                cost_cell = cell_state["Installed Cost"]
                is_nested = isinstance(cost_cell, str) and ":" in cost_cell
                tech_present = is_nested and re.search(
                    re.escape(tech) + r"\s*:", cost_cell, re.IGNORECASE)
                if not is_nested or tech_present:
                    old_cost = find_value_for_key(cost_cell, tech)
                    if is_zero_cost_placeholder(old_cost):
                        diff_rows.append({
                            "Name": row.get("Name"), "technology": tech,
                            "column": "Installed Cost",
                            "old_value": old_cost, "new_value": "(skipped)",
                            "btb_technology_id": entry["btb_technology_id"],
                            "btb_display_name": entry["btb_display_name"],
                            "projection_scenario":
                                entry["projection_scenario"],
                            "projection_year": entry["projection_year"],
                            "notes": "existing cost is $0 -- likely "
                                     "intentionally shared with a paired "
                                     "heating/cooling technology in this "
                                     "row to avoid double-counting; left "
                                     "unchanged, review manually",
                        })
                    else:
                        new_full_cost_cell = replace_cost_for_key(
                            cost_cell, tech, cost)
                        new_cost_repr = (
                            f"new: {cost['new']}; "
                            f"existing: {cost['existing']}")
                        if new_full_cost_cell != cost_cell:
                            cell_state["Installed Cost"] = new_full_cost_cell
                            ws.cell(
                                row=excel_row,
                                column=col_index["Installed Cost"]).value = \
                                new_full_cost_cell
                            diff_rows.append({
                                "Name": row.get("Name"), "technology": tech,
                                "column": "Installed Cost",
                                "old_value": old_cost,
                                "new_value": new_cost_repr,
                                "btb_technology_id":
                                    entry["btb_technology_id"],
                                "btb_display_name":
                                    entry["btb_display_name"],
                                "projection_scenario":
                                    entry["projection_scenario"],
                                "projection_year": entry["projection_year"],
                            })

            # -- Lifetime --
            lifetime = btb_row.get("Lifetime (Years)")
            if pd.notna(lifetime):
                lifetime_cell = cell_state["Lifetime"]
                old_lifetime = find_value_for_key(lifetime_cell, tech)
                new_lifetime_str = f"{float(lifetime):.3g}"
                if old_lifetime != new_lifetime_str and (
                        not isinstance(lifetime_cell, str)
                        or ":" not in lifetime_cell
                        or re.search(re.escape(tech) + r"\s*:",
                                     lifetime_cell, re.IGNORECASE)):
                    cell_state["Lifetime"] = replace_value_for_key(
                        lifetime_cell, tech, new_lifetime_str)
                    ws.cell(row=excel_row,
                            column=col_index["Lifetime"]).value = \
                        cell_state["Lifetime"]
                    diff_rows.append({
                        "Name": row.get("Name"), "technology": tech,
                        "column": "Lifetime",
                        "old_value": old_lifetime,
                        "new_value": new_lifetime_str,
                        "btb_technology_id": entry["btb_technology_id"],
                        "btb_display_name": entry["btb_display_name"],
                        "projection_scenario": entry["projection_scenario"],
                        "projection_year": entry["projection_year"],
                    })

    if "BTB Diff" in wb.sheetnames:
        del wb["BTB Diff"]
    diff_ws = wb.create_sheet("BTB Diff")
    diff_columns = [
        "Name", "technology", "column", "old_value", "new_value",
        "btb_technology_id", "btb_display_name", "projection_scenario",
        "projection_year", "notes"]
    diff_df = pd.DataFrame(diff_rows, columns=diff_columns).fillna("")
    diff_ws.append(diff_columns)
    for _, diff_row in diff_df.iterrows():
        diff_ws.append(list(diff_row))

    wb.save(OUT_PATH)
    print(f"Proposed {len(diff_rows)} cell changes across "
          f"{diff_df['Name'].nunique() if len(diff_df) else 0} measures.")
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
